import streamlit as st
import re
import pandas as pd
import io
from collections import defaultdict
from fpdf import FPDF
import random

# ===================== PAGE CONFIG =====================
st.set_page_config(page_title="iPa Result Analyzer v7.2", layout="wide", initial_sidebar_state="expanded")

# ===================== CSS =====================
st.markdown("""
<style>
    .stApp { background-color: #f1f5f9; }
    section[data-testid="stSidebar"] { background-color: #0f172a !important; padding-top: 2rem; }
    section[data-testid="stSidebar"] * { color: #e2e8f0 !important; }
    section[data-testid="stSidebar"] .stMarkdown h1, section[data-testid="stSidebar"] .stMarkdown h2 { color: #f59e0b !important; }
    div[data-testid="stMetric"] { background-color: #ffffff !important; padding: 15px !important; border-radius: 16px !important; box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.1) !important; border-left: 6px solid #f59e0b !important; }
    div[data-testid="stMetric"] p { color: #1e293b !important; font-weight: 600 !important; }
    h1, h2, h3 { color: #0f172a !important; font-weight: 700 !important; }
    .stFileUploader { background-color: white; border: 2px dashed #cbd5e1; border-radius: 16px; padding: 20px; }
</style>
""", unsafe_allow_html=True)

st.sidebar.markdown("## 🧠 iPa Analyzer v7.2")
st.sidebar.markdown("---")
st.sidebar.markdown("**📦 Features:**")
st.sidebar.markdown("- UGC Credit-Based SGPA (26 Credits)")
st.sidebar.markdown("- Export Excel & PDF")
st.sidebar.markdown("- Tie-Breaker + Smart Space Injector")
st.sidebar.markdown("- Enrollment No. Column")
st.sidebar.markdown("---")
st.sidebar.caption("Built with ❤️ for Moid | Deen + Dunya")

uploaded_file = st.file_uploader(
    "Upload Document (TXT, PDF, Image, CSV, Excel)",
    type=['csv', 'xlsx', 'pdf', 'png', 'jpg', 'jpeg', 'txt']
)

# ======================== SMART SPACE INJECTOR (FIXED) ========================
def clean_glued_name(name):
    """Convert glued names (MOIDAHMAD) into spaced names (MOID AHMAD)"""
    if not name:
        return name
    
    # Already has spaces — just trim and return
    if ' ' in name:
        return ' '.join(name.split())
    
    # Comprehensive name tokens (sorted by length descending)
    tokens = [
        'MOHAMMED', 'QUTUBUDDIN', 'MOHD', 'ABDUL', 'SHAIKH', 'SYED', 
        'MOHAMMAD', 'MOHAMED', 'SHAIK', 'ABDULLAH', 'NAEEM', 'AHMAD', 
        'JAVED', 'NAGMA', 'SHAFI', 'ASRAR', 'MEER', 'YOUSUF', 
        'HAROON', 'RAHMAN', 'WAHAB', 'FARID', 'ABBAS', 'SALMAN', 
        'AKHTAR', 'ANJUM', 'SABIR', 'DANISH', 'JUNAID', 'SIMRAN', 
        'KAUSAR', 'ABUZAR', 'MAQDOOM', 'IFTIKHAR', 'ZAMEER', 
        'ALAM', 'TABBASUM', 'PERVEEN', 'RAZA', 'HUSSAIN', 
        'NOMAN', 'ARIF', 'FAIZ', 'MEHAR', 'NISAR', 'SHAMS', 
        'AFTAB', 'SAFWAN', 'TAUQEER', 'SAQIB', 'BANO', 'MUGERA', 
        'ADNAN', 'KAMAL', 'AHMED', 'ASIF', 'ANWER', 'SIDDIQUI', 
        'FARAZ', 'KASHIF', 'NASEEM', 'RIZWAN', 'EHTASHAM', 
        'ISHFAQ', 'MAQBOOL', 'NAYYAR', 'SOHAIL', 'TANVEER', 
        'ZAHID', 'SHAHID', 'SHAKIR', 'MOID', 'AHMAD', 'ALI', 
        'KHAN', 'BIN', 'MD', 'HASAN', 'ABD', 'RAHMAT', 'HADI',
        'UL', 'WARA', 'BEGUM', 'FATIMA', 'ZEESHAN', 'FURQAN'
    ]
    
    # Sort by length descending (longer tokens first)
    tokens_sorted = sorted(tokens, key=len, reverse=True)
    
    result = name
    for token in tokens_sorted:
        new_result = ''
        i = 0
        while i < len(result):
            # Check if token matches at this position
            if i + len(token) <= len(result) and result[i:i+len(token)] == token:
                # If token is not at start and previous char is a letter (glued case)
                if i > 0 and result[i-1].isalpha():
                    new_result += ' ' + token
                else:
                    new_result += token
                i += len(token)
            else:
                new_result += result[i]
                i += 1
        result = new_result
    
    # Clean up extra spaces
    result = ' '.join(result.split())
    return result

# ======================== CORE PARSER ========================
def parse_mba_odl(text):
    lines = text.split('\n')
    data = []
    
    credits = [3, 3, 3, 3, 3, 3, 4, 2, 2]
    total_credits = 26
    
    for line in lines:
        parts = line.split()
        if len(parts) < 20:
            continue
            
        try:
            res_idx = parts.index('PASSED') if 'PASSED' in parts else parts.index('FAILED')
        except ValueError:
            continue
        
        sub_start = -1
        for i in range(4, res_idx):
            token = parts[i]
            if token.isdigit() or token == 'Ab':
                sub_start = i
                break
        
        if sub_start == -1:
            continue
        
        # --- BASIC INFO ---
        sl = parts[0]
        rc = parts[1]
        enrl_no = parts[2]
        
        # --- ROLL NO + NAME ---
        raw_roll = parts[3]
        roll_match = re.match(r'^(25MBAQ\d{3}[A-Z]{2})(.*)', raw_roll)
        
        if roll_match:
            roll = roll_match.group(1)
            name_extra = roll_match.group(2).strip()
        else:
            roll = raw_roll
            name_extra = ""
        
        name_tokens = parts[4:sub_start]
        if name_extra:
            full_name = name_extra
        else:
            full_name = ' '.join(name_tokens)
        
        # --- SMART SPACE INJECTOR (FIXED) ---
        full_name = clean_glued_name(full_name)
        if not full_name:
            full_name = roll
        
        # --- SUBJECT EXTRACTION ---
        subject_tokens = parts[sub_start:res_idx]
        if len(subject_tokens) < 25:
            continue
        
        weighted_sum = 0
        raw_total = 0
        
        for i in range(7):
            a_str = subject_tokens[i*3]
            t_str = subject_tokens[i*3 + 1]
            a = float(re.sub(r'[^0-9.]', '', a_str)) if a_str.replace('.','',1).isdigit() or a_str.isdigit() else 0
            t = float(re.sub(r'[^0-9.]', '', t_str)) if t_str.replace('.','',1).isdigit() or t_str.isdigit() else 0
            marks = a + t
            raw_total += marks
            
            if marks >= 90: gp = 10
            elif marks >= 75: gp = 9
            elif marks >= 60: gp = 8
            elif marks >= 55: gp = 7
            elif marks >= 50: gp = 6
            elif marks >= 45: gp = 5
            elif marks >= 40: gp = 4
            else: gp = 0
            weighted_sum += gp * credits[i]
        
        p8 = float(re.sub(r'[^0-9.]', '', subject_tokens[21])) if subject_tokens[21].replace('.','',1).isdigit() or subject_tokens[21].isdigit() else 0
        raw_total += p8
        pct_ls = (p8 / 50) * 100
        if pct_ls >= 90: gp_ls = 10
        elif pct_ls >= 75: gp_ls = 9
        elif pct_ls >= 60: gp_ls = 8
        elif pct_ls >= 55: gp_ls = 7
        elif pct_ls >= 50: gp_ls = 6
        elif pct_ls >= 45: gp_ls = 5
        elif pct_ls >= 40: gp_ls = 4
        else: gp_ls = 0
        weighted_sum += gp_ls * credits[7]
        
        p9 = float(re.sub(r'[^0-9.]', '', subject_tokens[23])) if subject_tokens[23].replace('.','',1).isdigit() or subject_tokens[23].isdigit() else 0
        raw_total += p9
        pct_ea = (p9 / 50) * 100
        if pct_ea >= 90: gp_ea = 10
        elif pct_ea >= 75: gp_ea = 9
        elif pct_ea >= 60: gp_ea = 8
        elif pct_ea >= 55: gp_ea = 7
        elif pct_ea >= 50: gp_ea = 6
        elif pct_ea >= 45: gp_ea = 5
        elif pct_ea >= 40: gp_ea = 4
        else: gp_ea = 0
        weighted_sum += gp_ea * credits[8]
        
        sgpa = round(weighted_sum / total_credits, 2)
        
        if sgpa >= 9.0: grade = 'O'
        elif sgpa >= 7.5: grade = 'A+'
        elif sgpa >= 6.0: grade = 'A'
        elif sgpa >= 5.5: grade = 'B+'
        elif sgpa >= 5.0: grade = 'B'
        elif sgpa >= 4.5: grade = 'C'
        elif sgpa >= 4.0: grade = 'D'
        else: grade = 'F'
        
        final_result = 'PASSED' if 'PASSED' in parts[res_idx] else 'FAILED'
        
        data.append({
            "SL": sl,
            "RC/SRC": rc,
            "Enrollment No": enrl_no,
            "Roll No": roll,
            "Student Name": full_name,
            "Total": int(raw_total),
            "SGPA": sgpa,
            "Grade": grade,
            "Result": final_result
        })
    
    if not data:
        return pd.DataFrame()
    
    df = pd.DataFrame(data)
    df = df.sort_values(by=['SGPA', 'Total', 'Student Name'], ascending=[False, False, True])
    df['Rank'] = range(1, len(df) + 1)
    df = df[['Rank', 'SL', 'Student Name', 'Enrollment No', 'Roll No', 'Total', 'SGPA', 'Grade', 'Result']]
    return df

# ======================== PDF GENERATOR ========================
def generate_pdf(df):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    
    pdf.set_font("Arial", 'B', 14)
    pdf.cell(200, 10, txt="MBA ODL Result Analysis (UGC Credit Rules)", ln=True, align='C')
    pdf.ln(8)
    
    pdf.set_font("Arial", 'B', 8)
    headers = ['Rank', 'Student Name', 'Enrollment', 'Roll No', 'Total', 'SGPA', 'Grade', 'Result']
    col_widths = [12, 50, 22, 25, 15, 15, 12, 18]
    for i, h in enumerate(headers):
        pdf.cell(col_widths[i], 10, h, border=1, align='C')
    pdf.ln()
    
    pdf.set_font("Arial", '', 7)
    for _, row in df.iterrows():
        pdf.cell(col_widths[0], 8, str(row['Rank']), border=1, align='C')
        pdf.cell(col_widths[1], 8, str(row['Student Name'])[:22], border=1, align='L')
        pdf.cell(col_widths[2], 8, str(row['Enrollment No']), border=1, align='C')
        pdf.cell(col_widths[3], 8, str(row['Roll No']), border=1, align='C')
        pdf.cell(col_widths[4], 8, str(row['Total']), border=1, align='C')
        pdf.cell(col_widths[5], 8, f"{row['SGPA']:.2f}", border=1, align='C')
        pdf.cell(col_widths[6], 8, str(row['Grade']), border=1, align='C')
        pdf.cell(col_widths[7], 8, str(row['Result']), border=1, align='C')
        pdf.ln()
    
    pdf.ln(6)
    pdf.set_font("Arial", 'I', 9)
    pdf.cell(200, 10, txt="Generated by iPa v7.2 | For Moid", ln=True, align='C')
    
    return pdf.output(dest='S').encode('latin1')

# ======================== MAIN ENGINE ========================

if uploaded_file is not None:
    file_name = uploaded_file.name.lower()
    df_result = pd.DataFrame()
    
    with st.spinner("🔄 iPa v7.2 analyzing..."):
        try:
            if file_name.endswith('.txt'):
                content = uploaded_file.read().decode('utf-8')
                df_result = parse_mba_odl(content)
                st.success("✅ Raw Text Data Parsed Successfully!")
            
            elif file_name.endswith('.csv'):
                df = pd.read_csv(uploaded_file)
                if len(df.columns) > 15:
                    txt_buffer = io.StringIO()
                    df.to_csv(txt_buffer, index=False, header=False)
                    df_result = parse_mba_odl(txt_buffer.getvalue())
                else:
                    df_result = df
                    st.info("📊 CSV uploaded. Displaying as is.")
            
            elif file_name.endswith('.xlsx'):
                import openpyxl
                df = pd.read_excel(uploaded_file)
                if len(df.columns) > 15:
                    txt_buffer = io.StringIO()
                    df.to_csv(txt_buffer, index=False, header=False)
                    df_result = parse_mba_odl(txt_buffer.getvalue())
                else:
                    df_result = df
                    st.info("📊 Excel uploaded. Displaying as is.")
            
            elif file_name.endswith(('pdf', 'png', 'jpg', 'jpeg')):
                text = ""
                if file_name.endswith('.pdf'):
                    try:
                        import pdfplumber
                        with pdfplumber.open(uploaded_file) as pdf:
                            for page in pdf.pages:
                                text += page.extract_text() or ""
                    except ImportError:
                        st.error("❌ pdfplumber not installed.")
                        st.stop()
                else:
                    try:
                        import pytesseract
                        from PIL import Image
                        image = Image.open(uploaded_file)
                        text = pytesseract.image_to_string(image)
                    except ImportError:
                        st.error("❌ pytesseract not installed.")
                        st.stop()
                
                if text:
                    df_result = parse_mba_odl(text)
                    st.success("✅ OCR/PDF Data Parsed Successfully!")
                else:
                    st.error("❌ No text extracted.")
                    st.stop()
            else:
                st.error("❌ Unsupported format")
                st.stop()

            if not df_result.empty:
                total_students = len(df_result)
                avg_sgpa = df_result['SGPA'].mean()
                pass_count = df_result[df_result['Result'] == 'PASSED'].shape[0]
                fail_count = total_students - pass_count
                
                col1, col2, col3, col4 = st.columns(4)
                col1.metric("👨‍🎓 Total", total_students)
                col2.metric("📈 Avg SGPA", f"{avg_sgpa:.2f}/10")
                col3.metric("✅ Passed", pass_count)
                col4.metric("❌ Failed", fail_count)
                
                st.markdown("---")
                
                st.subheader("🏆 Top 3 Toppers")
                top3 = df_result.head(3)
                for _, row in top3.iterrows():
                    st.markdown(f"**Rank {row['Rank']}:** {row['Student Name']} | SGPA: {row['SGPA']} | Total: {row['Total']} | Grade: {row['Grade']}")
                
                st.markdown("---")
                
                st.subheader("📋 Complete Ranked List")
                display_df = df_result.copy()
                display_df.index = range(1, len(display_df) + 1)
                st.dataframe(display_df.style.hide(axis='index'), use_container_width=True, height=500)
                
                st.markdown("---")
                st.subheader("📤 Export Your Report")
                
                col1, col2 = st.columns(2)
                
                with col1:
                    towrite = io.BytesIO()
                    with pd.ExcelWriter(towrite, engine='xlsxwriter') as writer:
                        df_result.to_excel(writer, sheet_name='Rank_List', index=False)
                    towrite.seek(0)
                    st.download_button(
                        label="📥 Download Excel (.xlsx)",
                        data=towrite,
                        file_name="MBA_Rank_List.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
                
                with col2:
                    try:
                        pdf_bytes = generate_pdf(df_result)
                        st.download_button(
                            label="📄 Download PDF Report",
                            data=pdf_bytes,
                            file_name="MBA_Rank_List.pdf",
                            mime="application/pdf"
                        )
                    except Exception as e:
                        st.error(f"PDF Error: {e}. pip install fpdf")
                
            else:
                st.warning("⚠️ No valid data found.")

        except Exception as e:
            st.error(f"🔥 Engine Error: {e}")
            st.info("💡 If you have raw text, paste it into a .txt file and upload.")

else:
    st.info("👆 Upload your MBA ODL Result data.")

st.markdown("---")
st.markdown("<p style='text-align: center; color: #64748b;'>Built with ❤️ by iPa v7.2 | Smart Space Injector | Rank Only</p>", unsafe_allow_html=True)
